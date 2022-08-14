import os
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

try:
    if not os.path.exists('out'):
        os.mkdir('out')
except Exception as e:
    print(e)
    
wb = load_workbook('./out/sample.xlsx') # 파일에서 워크북 불러옴
ws = wb.active

# B2:C11 까지의 데이터를 차트로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()
# bar_chart.add_data(bar_value)

# ws.add_chart(bar_chart, 'E1') # 차트 넣을 위치 정의

# B1:C11 까지의 데이터를 차트로 생성
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
# 계열 영어, 수학
line_chart.add_data(line_value, titles_from_data=True)
line_chart.title = '성적표'
line_chart.style = 10 # 미리 정의된 스타일을 적용, 사용자가 개별 지정도 가능
line_chart.x_axis.title = '번호'
line_chart.y_axis.title = '점수'

ws.add_chart(line_chart, 'E1')

wb.save('./out/sample_chart.xlsx')
